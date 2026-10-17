"""Unit tests for CSV and Excel parsers, XML security, and formula detection.

Governed by DATA_AND_API.md Section 4, TEST_AND_RELEASE_PLAN.md TM-08 & TM-13.2.
"""

import io
import zipfile

import pytest

from opwise.domain.exceptions import DataExchangeError, SecurityError
from opwise.io.csv_parser import parse_csv_bundle, parse_csv_zip_archive
from opwise.io.excel_parser import parse_excel_workbook
from opwise.io.security import inspect_safe_zip, is_formula_cell, validate_file_size
from opwise.io.templates import (
    generate_csv_template_zip,
    generate_excel_template_workbook,
)


def _create_test_zip(files: dict[str, str]) -> bytes:
    """Helper to construct a zip archive in memory."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for fname, content in files.items():
            zf.writestr(fname, content)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Security & Pre-parsing Tests
# ---------------------------------------------------------------------------


def test_file_size_validation() -> None:
    """Enforce non-empty and max file size boundaries."""
    with pytest.raises(DataExchangeError, match="empty"):
        validate_file_size(b"")

    validate_file_size(b"a" * 1024)

    with pytest.raises(DataExchangeError, match="exceeds the maximum"):
        validate_file_size(b"a" * (10 * 1024 * 1024 + 1))


def test_zip_path_traversal_rejection() -> None:
    """Reject ZIP archives containing path traversal sequences."""
    zip_bytes = _create_test_zip({"../evil.csv": "col1,col2\n1,2"})
    with pytest.raises(SecurityError, match="Potentially unsafe relative or absolute path"):
        inspect_safe_zip(zip_bytes)

    zip_bytes_abs = _create_test_zip({"/etc/passwd": "root:x:0:0"})
    with pytest.raises(SecurityError, match="Potentially unsafe relative or absolute path"):
        inspect_safe_zip(zip_bytes_abs)


def test_xxe_entity_injection_rejection() -> None:
    """Reject archives containing malicious XML DTD / ENTITY definitions (XXE / Billion Laughs)."""
    malicious_xml = b'<?xml version="1.0"?><!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><foo>&xxe;</foo>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", malicious_xml)
    zip_bytes = buf.getvalue()

    with pytest.raises(SecurityError, match="Malicious XML DTD or entity declaration detected"):
        inspect_safe_zip(zip_bytes, is_xlsx=True)


def test_formula_detection() -> None:
    """Identify spreadsheet formula injection patterns while allowing valid numbers."""
    assert is_formula_cell("=SUM(A1:A5)")
    assert is_formula_cell("@HYPERLINK('http://evil.com')")
    assert is_formula_cell("=1+1")
    assert is_formula_cell("  =cmd|' /C calc'!A0")
    assert is_formula_cell("+cmd|' /C calc'!A0")
    assert is_formula_cell("\t=test")

    # Valid literal numbers with +/- signs must NOT be flagged as formulas
    assert not is_formula_cell("+10")
    assert not is_formula_cell("-5")
    assert not is_formula_cell("-3.14")
    assert not is_formula_cell("Normal Text")
    assert not is_formula_cell(100)
    assert not is_formula_cell(None)


# ---------------------------------------------------------------------------
# CSV Parser Unit Tests
# ---------------------------------------------------------------------------


def test_valid_csv_template_round_trip() -> None:
    """Starter CSV template parses completely free of errors."""
    zip_bytes = generate_csv_template_zip()
    parsed = parse_csv_zip_archive(zip_bytes)
    assert parsed.is_valid
    assert len(parsed.errors) == 0
    assert len(parsed.machines) == 3
    assert len(parsed.orders) == 3
    assert len(parsed.operations) == 6
    assert len(parsed.downtimes) == 2


def test_csv_missing_required_files() -> None:
    """Reject bundle missing required files (e.g. operations.csv)."""
    zip_bytes = _create_test_zip({
        "machines.csv": "machine_id,name\nM1,Center 1",
        "orders.csv": "order_id,name,release_minute,due_minute,tardiness_weight\nORD-1,Order 1,0,10,1",
    })
    parsed = parse_csv_zip_archive(zip_bytes)
    assert not parsed.is_valid
    assert any("operations.csv" in err.sheet for err in parsed.errors)


def test_csv_missing_columns() -> None:
    """Report row error when expected column header is missing."""
    file_map = {
        "machines.csv": "machine_id\nM1",  # missing 'name'
        "orders.csv": "order_id,name,release_minute,due_minute,tardiness_weight\nORD-1,Order 1,0,10,1",
        "operations.csv": "operation_id,order_id,sequence_index,machine_id,processing_minutes\nOP-1,ORD-1,1,M1,10",
    }
    parsed = parse_csv_bundle(file_map)
    assert not parsed.is_valid
    assert any("Missing required column" in err.issue and "name" in err.issue for err in parsed.errors)


def test_csv_invalid_types_and_bounds() -> None:
    """Collect line-numbered errors for non-integer types and bound violations."""
    file_map = {
        "machines.csv": "machine_id,name\nM1,Center 1",
        "orders.csv": (
            "order_id,name,release_minute,due_minute,tardiness_weight\n"
            "ORD-1,Order 1,-5,10,1\n"  # negative release_minute
            "ORD-2,Order 2,0,10,0\n"   # weight < 1
            "ORD-3,Order 3,not_a_number,10,1\n"  # string in int column
        ),
        "operations.csv": (
            "operation_id,order_id,sequence_index,machine_id,processing_minutes\n"
            "OP-1,ORD-1,1,M1,0\n"  # duration <= 0
        ),
        "downtimes.csv": (
            "downtime_id,machine_id,start_minute,end_minute,reason\n"
            "DT-1,M1,50,30,Maintenance\n"  # end <= start
        ),
    }
    parsed = parse_csv_bundle(file_map)
    assert not parsed.is_valid
    assert len(parsed.errors) >= 5

    err_msgs = [e.issue for e in parsed.errors]
    assert any("violates minimum bound" in m for m in err_msgs)
    assert any("Must be an integer" in m for m in err_msgs)
    assert any("strictly greater than start" in m for m in err_msgs)


def test_csv_formula_cell_rejected() -> None:
    """Formula cell in CSV is caught and reported as error."""
    file_map = {
        "machines.csv": "machine_id,name\n=CMD|'calc'!A0,Center 1",
        "orders.csv": "order_id,name,release_minute,due_minute,tardiness_weight\nORD-1,Order 1,=SUM(1,2),10,1",
        "operations.csv": "operation_id,order_id,sequence_index,machine_id,processing_minutes\nOP-1,ORD-1,1,M1,10",
    }
    parsed = parse_csv_bundle(file_map)
    assert not parsed.is_valid
    assert any("formulas are not permitted" in e.issue for e in parsed.errors)


# ---------------------------------------------------------------------------
# Excel Parser Unit Tests
# ---------------------------------------------------------------------------


def test_valid_excel_template_round_trip() -> None:
    """Starter Excel template parses cleanly with all sheets."""
    xlsx_bytes = generate_excel_template_workbook()
    parsed = parse_excel_workbook(xlsx_bytes)
    assert parsed.is_valid
    assert len(parsed.errors) == 0
    assert len(parsed.machines) == 3
    assert len(parsed.orders) == 3
    assert len(parsed.operations) == 6
    assert len(parsed.downtimes) == 2


def test_excel_missing_sheet() -> None:
    """Reject workbook missing required sheet (e.g. Operations)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Machines"
    ws_m.append(["machine_id", "name"])
    ws_m.append(["M1", "Center 1"])

    ws_o = wb.create_sheet("Orders")
    ws_o.append(["order_id", "name", "release_minute", "due_minute", "tardiness_weight"])
    ws_o.append(["ORD-1", "Order 1", 0, 10, 1])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    parsed = parse_excel_workbook(buf.getvalue())
    assert not parsed.is_valid
    assert any("Operations" in e.sheet and "missing required sheet" in e.issue for e in parsed.errors)


def test_excel_formula_in_cells() -> None:
    """Reject formula expressions embedded in Excel cells."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws_m = wb.active
    ws_m.title = "Machines"
    ws_m.append(["machine_id", "name"])
    ws_m.append(["M1", "Center 1"])

    ws_o = wb.create_sheet("Orders")
    ws_o.append(["order_id", "name", "release_minute", "due_minute", "tardiness_weight"])
    ws_o.append(["ORD-1", "=CONCAT('A', 'B')", 0, 10, 1])

    ws_op = wb.create_sheet("Operations")
    ws_op.append(["operation_id", "order_id", "sequence_index", "machine_id", "processing_minutes"])
    ws_op.append(["OP-1", "ORD-1", "=1+0", "M1", 10])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    parsed = parse_excel_workbook(buf.getvalue())
    assert not parsed.is_valid
    assert any("formulas are not permitted" in e.issue for e in parsed.errors)
