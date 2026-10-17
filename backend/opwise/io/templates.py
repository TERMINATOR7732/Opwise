"""Starter template generators for CSV bundles and multi-sheet Excel files.

Governed by DATA_AND_API.md Section 3.2 (GET /templates/{format}) and Section 4.
"""

import io
import zipfile

import openpyxl

SAMPLE_MACHINES_CSV = """machine_id,name
M1,CNC Milling Center 1
M2,Lathe Turning Station
M3,Surface Grinder
"""

SAMPLE_ORDERS_CSV = """order_id,name,release_minute,due_minute,tardiness_weight
ORD-101,Titanium Bracket A,0,120,5
ORD-102,Aluminum Housing B,15,180,2
ORD-103,Steel Shaft C,30,90,10
"""

SAMPLE_OPERATIONS_CSV = """operation_id,order_id,sequence_index,machine_id,processing_minutes
OP-101-1,ORD-101,1,M1,35
OP-101-2,ORD-101,2,M2,25
OP-102-1,ORD-102,1,M2,40
OP-102-2,ORD-102,2,M3,50
OP-103-1,ORD-103,1,M1,20
OP-103-2,ORD-103,2,M3,30
"""

SAMPLE_DOWNTIMES_CSV = """downtime_id,machine_id,start_minute,end_minute,reason
DT-01,M1,60,90,Spindle Calibration
DT-02,M2,120,150,Tooling Replacement
"""


def generate_csv_template_zip() -> bytes:
    """Generate a valid starter ZIP archive containing sample CSV files."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("machines.csv", SAMPLE_MACHINES_CSV)
        zf.writestr("orders.csv", SAMPLE_ORDERS_CSV)
        zf.writestr("operations.csv", SAMPLE_OPERATIONS_CSV)
        zf.writestr("downtimes.csv", SAMPLE_DOWNTIMES_CSV)
    return buffer.getvalue()


def generate_excel_template_workbook() -> bytes:
    """Generate a valid starter Excel workbook containing sample multi-sheet data."""
    wb = openpyxl.Workbook()
    # Default sheet rename to Machines
    ws_m = wb.active
    ws_m.title = "Machines"
    ws_m.append(["machine_id", "name"])
    ws_m.append(["M1", "CNC Milling Center 1"])
    ws_m.append(["M2", "Lathe Turning Station"])
    ws_m.append(["M3", "Surface Grinder"])

    ws_o = wb.create_sheet("Orders")
    ws_o.append(["order_id", "name", "release_minute", "due_minute", "tardiness_weight"])
    ws_o.append(["ORD-101", "Titanium Bracket A", 0, 120, 5])
    ws_o.append(["ORD-102", "Aluminum Housing B", 15, 180, 2])
    ws_o.append(["ORD-103", "Steel Shaft C", 30, 90, 10])

    ws_op = wb.create_sheet("Operations")
    ws_op.append(["operation_id", "order_id", "sequence_index", "machine_id", "processing_minutes"])
    ws_op.append(["OP-101-1", "ORD-101", 1, "M1", 35])
    ws_op.append(["OP-101-2", "ORD-101", 2, "M2", 25])
    ws_op.append(["OP-102-1", "ORD-102", 1, "M2", 40])
    ws_op.append(["OP-102-2", "ORD-102", 2, "M3", 50])
    ws_op.append(["OP-103-1", "ORD-103", 1, "M1", 20])
    ws_op.append(["OP-103-2", "ORD-103", 2, "M3", 30])

    ws_dt = wb.create_sheet("Downtimes")
    ws_dt.append(["downtime_id", "machine_id", "start_minute", "end_minute", "reason"])
    ws_dt.append(["DT-01", "M1", 60, 90, "Spindle Calibration"])
    ws_dt.append(["DT-02", "M2", 120, 150, "Tooling Replacement"])

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
