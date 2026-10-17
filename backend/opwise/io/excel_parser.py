"""OpenPyXL multi-sheet Excel (.xlsx) workbook parser and validator.

Governed by DATA_AND_API.md Section 4.2 and Section 4.3.
Parses multi-sheet workbooks (Machines, Orders, Operations, Downtimes)
with XML security validation, formula detection, and row-level error reporting.
"""

import io
from typing import Any

import openpyxl

from opwise.io.parsing_types import (
    ParsedDowntime,
    ParsedFactoryData,
    ParsedMachine,
    ParsedOperation,
    ParsedOrder,
    RowError,
)
from opwise.io.security import (
    MAX_ROW_COUNT,
    inspect_safe_zip,
    is_formula_cell,
)

REQUIRED_SHEETS = {"machines", "orders", "operations"}

MACHINE_COLS = ["machine_id", "name"]
ORDER_COLS = ["order_id", "name", "release_minute", "due_minute", "tardiness_weight"]
OP_COLS = ["operation_id", "order_id", "sequence_index", "machine_id", "processing_minutes"]
DT_COLS = ["downtime_id", "machine_id", "start_minute", "end_minute", "reason"]


def _find_sheet(wb: openpyxl.Workbook, target_name: str) -> Any | None:
    """Find a sheet in workbook by case-insensitive name."""
    target_lower = target_name.lower().strip()
    for name in wb.sheetnames:
        if name.lower().strip() == target_lower:
            return wb[name]
    return None


def _get_header_col_index(sheet: Any, expected_cols: list[str], sheet_name: str, errors: list[RowError]) -> dict[str, int] | None:
    """Map expected column names to 0-based column indices from row 1."""
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        errors.append(RowError(sheet_name, 1, "header", "", "Sheet is empty or missing header row."))
        return None

    col_index: dict[str, int] = {}
    for idx, cell_val in enumerate(header_row):
        if cell_val is not None:
            clean_name = str(cell_val).strip().lower()
            col_index[clean_name] = idx

    missing = [c for c in expected_cols if c not in col_index]
    if missing:
        errors.append(
            RowError(sheet_name, 1, "header", ", ".join(missing), f"Missing required column(s): {', '.join(missing)}.")
        )
        return None
    return col_index


def _parse_int_cell(
    row_vals: tuple[Any, ...], col_idx: int, field_name: str, sheet_name: str, row_num: int, errors: list[RowError], min_val: int | None = None
) -> int | None:
    """Validate and extract integer cell, recording RowError on failure."""
    if col_idx >= len(row_vals):
        errors.append(RowError(sheet_name, row_num, field_name, "", "Missing cell value."))
        return None
    val = row_vals[col_idx]
    if val is None or str(val).strip() == "":
        errors.append(RowError(sheet_name, row_num, field_name, "", "Required integer value cannot be empty."))
        return None
    if is_formula_cell(val):
        errors.append(RowError(sheet_name, row_num, field_name, str(val), "Spreadsheet formulas are not permitted."))
        return None
    try:
        iv = int(val)
    except (ValueError, TypeError):
        errors.append(RowError(sheet_name, row_num, field_name, str(val), f"Must be an integer, got '{val}'."))
        return None

    if min_val is not None and iv < min_val:
        errors.append(
            RowError(sheet_name, row_num, field_name, str(val), f"Value {iv} violates minimum bound (must be >= {min_val}).")
        )
        return None
    return iv


def _parse_str_cell(
    row_vals: tuple[Any, ...], col_idx: int, field_name: str, sheet_name: str, row_num: int, errors: list[RowError], required: bool = True
) -> str | None:
    """Validate and extract text cell, rejecting formulas."""
    if col_idx >= len(row_vals):
        if required:
            errors.append(RowError(sheet_name, row_num, field_name, "", "Missing cell value."))
        return None
    val = row_vals[col_idx]
    if val is None:
        if required:
            errors.append(RowError(sheet_name, row_num, field_name, "", "Required text value cannot be empty."))
        return None
    s_val = str(val).strip()
    if not s_val and required:
        errors.append(RowError(sheet_name, row_num, field_name, s_val, "Required text value cannot be empty."))
        return None
    if is_formula_cell(val):
        errors.append(RowError(sheet_name, row_num, field_name, str(val), "Spreadsheet formulas are not permitted."))
        return None
    return s_val


def _parse_machines_sheet(sheet: Any, sheet_name: str, errors: list[RowError]) -> list[ParsedMachine]:
    """Parse Machines worksheet."""
    col_map = _get_header_col_index(sheet, MACHINE_COLS, sheet_name, errors)
    if col_map is None:
        return []
    machines: list[ParsedMachine] = []
    for row_num, row_vals in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row_vals):
            continue  # Skip empty row
        if row_num > MAX_ROW_COUNT + 1:
            errors.append(RowError(sheet_name, row_num, "rows", "", f"Exceeded maximum row limit of {MAX_ROW_COUNT}."))
            break
        m_id = _parse_str_cell(row_vals, col_map["machine_id"], "machine_id", sheet_name, row_num, errors, required=True)
        name = _parse_str_cell(row_vals, col_map["name"], "name", sheet_name, row_num, errors, required=False)
        if m_id is not None:
            machines.append(ParsedMachine(machine_id=m_id, name=name or m_id))
    return machines


def _parse_orders_sheet(sheet: Any, sheet_name: str, errors: list[RowError]) -> list[ParsedOrder]:
    """Parse Orders worksheet."""
    col_map = _get_header_col_index(sheet, ORDER_COLS, sheet_name, errors)
    if col_map is None:
        return []
    orders: list[ParsedOrder] = []
    for row_num, row_vals in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row_vals):
            continue
        if row_num > MAX_ROW_COUNT + 1:
            errors.append(RowError(sheet_name, row_num, "rows", "", f"Exceeded maximum row limit of {MAX_ROW_COUNT}."))
            break
        o_id = _parse_str_cell(row_vals, col_map["order_id"], "order_id", sheet_name, row_num, errors, required=True)
        name = _parse_str_cell(row_vals, col_map["name"], "name", sheet_name, row_num, errors, required=False)
        rel = _parse_int_cell(row_vals, col_map["release_minute"], "release_minute", sheet_name, row_num, errors, min_val=0)
        due = _parse_int_cell(row_vals, col_map["due_minute"], "due_minute", sheet_name, row_num, errors, min_val=0)
        wt = _parse_int_cell(row_vals, col_map["tardiness_weight"], "tardiness_weight", sheet_name, row_num, errors, min_val=1)
        if o_id is not None and rel is not None and due is not None and wt is not None:
            orders.append(ParsedOrder(order_id=o_id, name=name or o_id, release_minute=rel, due_minute=due, tardiness_weight=wt))
    return orders


def _parse_operations_sheet(sheet: Any, sheet_name: str, errors: list[RowError]) -> list[ParsedOperation]:
    """Parse Operations worksheet."""
    col_map = _get_header_col_index(sheet, OP_COLS, sheet_name, errors)
    if col_map is None:
        return []
    ops: list[ParsedOperation] = []
    for row_num, row_vals in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row_vals):
            continue
        if row_num > MAX_ROW_COUNT + 1:
            errors.append(RowError(sheet_name, row_num, "rows", "", f"Exceeded maximum row limit of {MAX_ROW_COUNT}."))
            break
        op_id = _parse_str_cell(row_vals, col_map["operation_id"], "operation_id", sheet_name, row_num, errors, required=True)
        order_id = _parse_str_cell(row_vals, col_map["order_id"], "order_id", sheet_name, row_num, errors, required=True)
        machine_id = _parse_str_cell(row_vals, col_map["machine_id"], "machine_id", sheet_name, row_num, errors, required=True)
        seq = _parse_int_cell(row_vals, col_map["sequence_index"], "sequence_index", sheet_name, row_num, errors, min_val=1)
        dur = _parse_int_cell(row_vals, col_map["processing_minutes"], "processing_minutes", sheet_name, row_num, errors, min_val=1)
        if op_id is not None and order_id is not None and machine_id is not None and seq is not None and dur is not None:
            ops.append(ParsedOperation(operation_id=op_id, order_id=order_id, sequence_index=seq, machine_id=machine_id, processing_minutes=dur))
    return ops


def _parse_downtimes_sheet(sheet: Any, sheet_name: str, errors: list[RowError]) -> list[ParsedDowntime]:
    """Parse optional Downtimes worksheet."""
    col_map = _get_header_col_index(sheet, DT_COLS, sheet_name, errors)
    if col_map is None:
        return []
    dts: list[ParsedDowntime] = []
    for row_num, row_vals in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row_vals):
            continue
        if row_num > MAX_ROW_COUNT + 1:
            errors.append(RowError(sheet_name, row_num, "rows", "", f"Exceeded maximum row limit of {MAX_ROW_COUNT}."))
            break
        dt_id = _parse_str_cell(row_vals, col_map["downtime_id"], "downtime_id", sheet_name, row_num, errors, required=True)
        m_id = _parse_str_cell(row_vals, col_map["machine_id"], "machine_id", sheet_name, row_num, errors, required=True)
        reason = _parse_str_cell(row_vals, col_map["reason"], "reason", sheet_name, row_num, errors, required=False)
        start = _parse_int_cell(row_vals, col_map["start_minute"], "start_minute", sheet_name, row_num, errors, min_val=0)
        end = _parse_int_cell(row_vals, col_map["end_minute"], "end_minute", sheet_name, row_num, errors, min_val=0)
        if dt_id is not None and m_id is not None and start is not None and end is not None:
            if end <= start:
                errors.append(
                    RowError(sheet_name, row_num, "end_minute", str(end), f"End minute ({end}) must be strictly greater than start ({start}).")
                )
            else:
                dts.append(ParsedDowntime(downtime_id=dt_id, machine_id=m_id, start_minute=start, end_minute=end, reason=reason or ""))
    return dts


def parse_excel_workbook(file_bytes: bytes) -> ParsedFactoryData:
    """Inspect and parse an Excel workbook with multi-sheet validation and XML security."""
    inspect_safe_zip(file_bytes, is_xlsx=True)
    result = ParsedFactoryData()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
    except Exception as err:
        result.errors.append(RowError("Workbook", 0, "file", "", f"Failed to parse Excel workbook: {err}"))
        return result

    try:
        # Check required sheets presence
        available_lower = {s.lower().strip() for s in wb.sheetnames}
        missing_sheets = [s.capitalize() for s in REQUIRED_SHEETS if s not in available_lower]
        if missing_sheets:
            for ms in missing_sheets:
                result.errors.append(RowError(ms, 0, "sheet", "", f"Workbook missing required sheet: '{ms}'."))
            return result

        m_sheet = _find_sheet(wb, "Machines")
        if m_sheet is not None:
            result.machines = _parse_machines_sheet(m_sheet, "Machines", result.errors)

        o_sheet = _find_sheet(wb, "Orders")
        if o_sheet is not None:
            result.orders = _parse_orders_sheet(o_sheet, "Orders", result.errors)

        op_sheet = _find_sheet(wb, "Operations")
        if op_sheet is not None:
            result.operations = _parse_operations_sheet(op_sheet, "Operations", result.errors)

        dt_sheet = _find_sheet(wb, "Downtimes")
        if dt_sheet is not None:
            result.downtimes = _parse_downtimes_sheet(dt_sheet, "Downtimes", result.errors)
    finally:
        wb.close()

    return result
